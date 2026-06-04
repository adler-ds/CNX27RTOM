#!/usr/bin/env python3
"""
Script to extract login data from Excel file (RTOM tab) and save to JSON.
Run this if you need to refresh the logins data from the Excel file.

Usage:
    python3 scripts/extract-logins.py path/to/excel-file.xlsx
"""

import sys
import json
import openpyxl
from pathlib import Path


def extract_logins(excel_path: str, output_path: str = 'data/logins.json'):
    """Extract login data from RTOM tab of Excel file."""

    try:
        wb = openpyxl.load_workbook(excel_path)

        if 'RTOM' not in wb.sheetnames:
            print(f"Error: 'RTOM' sheet not found in {excel_path}")
            print(f"Available sheets: {wb.sheetnames}")
            sys.exit(1)

        ws = wb['RTOM']
        logins = []

        # Start from row 2 (skip header row which is empty in this case)
        for i in range(2, ws.max_row + 1):
            url = ws.cell(row=i, column=1).value
            username = ws.cell(row=i, column=2).value
            password = ws.cell(row=i, column=4).value

            # Only add rows with all three values present
            if url and username and password:
                logins.append({
                    'id': i - 2,  # 0-indexed ID
                    'url': str(url),
                    'username': str(username),
                    'password': str(password),
                    'assigned': False,
                    'assignedTo': None
                })

        # Ensure output directory exists
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        # Write to JSON file
        with open(output_path, 'w') as f:
            json.dump(logins, f, indent=2)

        print(f"✓ Successfully extracted {len(logins)} logins")
        print(f"✓ Saved to {output_path}")

    except FileNotFoundError:
        print(f"Error: File not found: {excel_path}")
        sys.exit(1)
    except Exception as e:
        print(f"Error: {str(e)}")
        sys.exit(1)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 scripts/extract-logins.py <excel-file-path>")
        print("\nExample:")
        print("  python3 scripts/extract-logins.py ../Connections\\ -\\ All\\ Orgs.xlsx")
        sys.exit(1)

    excel_file = sys.argv[1]
    extract_logins(excel_file)
